from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


NAVY = "173D93"
BLUE = "2459D3"
PALE_BLUE = "EAF1FF"
WHITE = "FFFFFF"
GREY = "5B6573"
LIGHT_GREY = "F3F5F8"
GREEN = "E2F0D9"
THIN_GREY = Side(style="thin", color="D9E0EA")


def build_excel_report(projection: dict[str, Any]) -> bytes:
    """Create an editable workbook whose assumptions drive the full annual forecast."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Executive Summary"
    assumptions = wb.create_sheet("Assumptions")
    inputs = wb.create_sheet("Model Inputs")
    costs = wb.create_sheet("Project Cost Breakdown")
    notes = wb.create_sheet("Read Me")

    name = projection.get("project_name", "New Project")
    project_type = projection.get("project_type", "Other")
    # This model calculates in INR only; do not label INR amounts with another currency.
    currency = "INR"
    annual = projection.get("annual_projection", [])
    model = projection.get("model_inputs", {})
    # Backward-compatible fallbacks derived from the supplied projection.
    first = annual[0] if annual else {}
    project_cost = projection.get("project_cost", 0) or 0
    debt = projection.get("debt", 0) or 0
    debt_pct = model.get("debt_pct", (debt / project_cost * 100 if project_cost else 50))
    monthly_revenue = model.get("monthly_revenue", (first.get("revenue", 0) or 0) / 12)
    growth_pct = model.get("growth_pct", 5)
    variable_pct = model.get("variable_cost_pct", ((first.get("variable_cost", 0) or 0) / first.get("revenue", 1) * 100 if first.get("revenue") else 45))
    monthly_fixed = model.get("monthly_fixed_cost", (first.get("fixed_cost", 0) or 0) / 12)
    interest_pct = model.get("interest_pct", ((first.get("interest", 0) or 0) / (first.get("opening_debt", 1) or 1) * 100 if first.get("opening_debt") else 11))
    # Keep legacy fallbacks inside the same bounds as the editable input validations.
    loan_years = model.get("loan_years", min(30, max(1, len(annual) or 7)))
    years = model.get("projection_years", min(100, max(1, len(annual) or 5)))

    inputs.append(["KAAMKITPRO FINANCE | EDITABLE MODEL INPUTS"])
    inputs.append(["Project name", name])
    inputs.append(["Project type", project_type])
    inputs.append(["Currency", currency])
    # User-controlled project details must remain literal text in Excel.
    # Apply this after assigning values so touching B4 cannot shift appended rows.
    for cell_ref in ("B2", "B3", "B4"):
        inputs[cell_ref].data_type = "s"
    inputs.append(["Project cost (calculated)", f'=IF(B14="Itemized costs",B16,B15)'])
    inputs.append(["Debt funding (%)", debt_pct])
    inputs.append(["Monthly revenue", monthly_revenue])
    inputs.append(["Annual revenue growth (%)", growth_pct])
    inputs.append(["Variable cost (%)", variable_pct])
    inputs.append(["Monthly fixed cost", monthly_fixed])
    inputs.append(["Annual interest rate (%)", interest_pct])
    inputs.append(["Loan term (years)", loan_years])
    inputs.append(["Projection period (years)", years])
    itemized_costs = model.get("itemized_costs")
    inputs.append(["Project cost basis", "Itemized costs" if itemized_costs is not None else "Aggregate input"])
    inputs.append(["Aggregate project cost", project_cost])
    inputs.append(["Itemized project costs total", "='Project Cost Breakdown'!B12"])
    inputs["B5"].number_format = '"INR " #,##0.00'
    inputs["B15"].number_format = '"INR " #,##0.00'
    inputs["B16"].number_format = '"INR " #,##0.00'
    for cell in ("B6", "B8", "B9", "B11"):
        inputs[cell].number_format = "0.00"
    for cell in ("B7", "B10"):
        inputs[cell].number_format = '"INR " #,##0.00'
    for cell in ("B6", "B7", "B8", "B9", "B10", "B11", "B12", "B13", "B14", "B15"):
        inputs[cell].fill = PatternFill("solid", fgColor="FFF2CC")
    # Guide spreadsheet users to valid assumptions before formulas recalculate.
    from openpyxl.worksheet.datavalidation import DataValidation

    validations = (
        ("decimal", "between", "0", "100", "B6", "Debt funding must be between 0% and 100%."),
        ("decimal", "between", "-100", "500", "B8", "Annual revenue growth must be between -100% and 500%."),
        ("decimal", "between", "0", "100", "B9", "Variable cost must be between 0% and 100%."),
        ("decimal", "between", "0", "100", "B11", "Annual interest rate must be between 0% and 100%."),
        ("decimal", "greaterThanOrEqual", "0", None, "B15", "Aggregate project cost cannot be negative."),
        ("decimal", "greaterThanOrEqual", "0", None, "B7", "Monthly revenue cannot be negative."),
        ("decimal", "greaterThanOrEqual", "0", None, "B10", "Monthly fixed cost cannot be negative."),
        ("whole", "between", "1", "30", "B12", "Loan term must be a whole number from 1 to 30 years."),
        ("whole", "between", "1", "100", "B13", "Projection period must be a whole number from 1 to 100 years."),
    )
    for kind, operator, formula1, formula2, cell, message in validations:
        validation = DataValidation(
            type=kind, operator=operator, formula1=formula1, formula2=formula2,
            allow_blank=False,
        )
        validation.errorTitle = "Invalid model input"
        validation.error = message
        validation.promptTitle = "Allowed input range"
        validation.prompt = message
        validation.showErrorMessage = True
        validation.showInputMessage = True
        inputs.add_data_validation(validation)
        validation.add(inputs[cell])
    basis_validation = DataValidation(type="list", formula1='"Aggregate input,Itemized costs"', allow_blank=False)
    basis_validation.errorTitle = "Select a cost basis"
    basis_validation.error = "Choose Aggregate input or Itemized costs."
    basis_validation.showErrorMessage = True
    inputs.add_data_validation(basis_validation)
    basis_validation.add(inputs["B14"])
    inputs["B5"].fill = PatternFill("solid", fgColor=GREEN)
    inputs["B16"].fill = PatternFill("solid", fgColor=LIGHT_GREY)
    inputs.column_dimensions["A"].width = 34
    inputs.column_dimensions["B"].width = 24
    inputs.freeze_panes = "A2"

    # Editable cost-planning worksheet. Select "Itemized costs" in Model Inputs
    # to use this total in the project funding and forecast calculations; the
    # default "Aggregate input" preserves existing workbook behaviour.
    costs.append(["KAAMKITPRO FINANCE | PROJECT COST BREAKDOWN"])
    costs.append(["Cost category", "Estimated cost (INR)", "Basis / note"])
    cost_rows = [
        ("Outsourcing", 1_500_000, "User-provided starting estimate"),
        ("Warehouse setup", 2_000_000, "User-provided minimum estimate"),
        ("Hospital setup / expansion", 0, "Enter estimate; zero means not yet specified"),
        ("Retail / pharmacy setup or expansion", 0, "Enter estimate; zero means not yet specified"),
        ("Optical setup or expansion", 0, "Enter estimate; zero means not yet specified"),
        ("Equipment and technology", 0, "Enter estimate; zero means not yet specified"),
        ("Licences, deposits and professional fees", 0, "Enter estimate; zero means not yet specified"),
        ("Initial inventory and working capital", 0, "Enter estimate; zero means not yet specified"),
        ("Other project costs", 0, "Enter estimate; zero means not yet specified"),
    ]
    if itemized_costs is not None:
        cost_rows = [
            (label, float(itemized_costs.get(label, 0)), note)
            for label, _default, note in cost_rows
        ]
    for label, amount, note in cost_rows:
        costs.append([label, amount, note])
    costs.append(["Total itemized project costs", "=SUM(B3:B11)", "Formula-driven total; reconcile before using in funding plan"])
    costs["B12"].number_format = '"INR " #,##0.00'
    for row in range(3, 12):
        costs.cell(row, 2).number_format = '"INR " #,##0.00'
        costs.cell(row, 2).fill = PatternFill("solid", fgColor="FFF2CC")
        costs.cell(row, 2).protection = costs.cell(row, 2).protection.copy(locked=False)
        costs.cell(row, 2).data_type = "n"
    costs["A12"].font = Font(bold=True)
    costs["B12"].font = Font(bold=True)
    costs["B12"].fill = PatternFill("solid", fgColor=GREEN)
    costs.column_dimensions["A"].width = 38
    costs.column_dimensions["B"].width = 24
    costs.column_dimensions["C"].width = 64
    costs.freeze_panes = "A3"
    costs.auto_filter.ref = "A2:C11"
    from openpyxl.worksheet.datavalidation import DataValidation
    cost_validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=False)
    cost_validation.errorTitle = "Invalid cost"
    cost_validation.error = "Enter a non-negative cost in INR."
    cost_validation.showErrorMessage = True
    costs.add_data_validation(cost_validation)
    cost_validation.add("B3:B11")

    # Detailed projection formulas refer to editable inputs; no fixed forecast values.
    assumptions.append(["KAAMKITPRO FINANCE | LIVE ANNUAL PROJECTION"])
    assumptions.append(["Project name", "='Model Inputs'!B2"])
    assumptions.append(["Project type", "='Model Inputs'!B3"])
    assumptions.append(["Currency", "='Model Inputs'!B4"])
    assumptions.append(["Project cost", "='Model Inputs'!B5"])
    assumptions.append(["Debt", "='Model Inputs'!B5*'Model Inputs'!B6/100"])
    assumptions.append(["Equity", "='Model Inputs'!B5-B6"])
    assumptions.append([])
    assumptions.append(["Year", "Opening Debt", "Revenue", "Variable Cost", "Fixed Cost", "EBITDA", "Interest", "Principal", "Debt Service", "DSCR", "Cash After Debt Service (Pre-Tax, Pre-Capex)", "Closing Debt", "Profit Before Tax (Pre-Depreciation)", "Break-even Revenue"])
    for year in range(1, 101):
        r = 9 + year
        assumptions.cell(r, 1, f'=IF(ROW()-9<=\'Model Inputs\'!$B$13,ROW()-9,"")')
        assumptions.cell(r, 2, f'=IF(A{r}="","",MAX(0,$B$6-(A{r}-1)*$B$6/\'Model Inputs\'!$B$12))')
        assumptions.cell(r, 3, f'=IF(A{r}="","",\'Model Inputs\'!$B$7*12*(1+\'Model Inputs\'!$B$8/100)^(A{r}-1))')
        assumptions.cell(r, 4, f'=IF(A{r}="","",C{r}*\'Model Inputs\'!$B$9/100)')
        assumptions.cell(r, 5, f'=IF(A{r}="","",\'Model Inputs\'!$B$10*12)')
        assumptions.cell(r, 6, f'=IF(A{r}="","",C{r}-D{r}-E{r})')
        assumptions.cell(r, 7, f'=IF(A{r}="","",IF(A{r}<=\'Model Inputs\'!$B$12,B{r}*\'Model Inputs\'!$B$11/100,0))')
        assumptions.cell(r, 8, f'=IF(A{r}="","",IF(A{r}<=\'Model Inputs\'!$B$12,MIN(B{r},$B$6/\'Model Inputs\'!$B$12),0))')
        assumptions.cell(r, 9, f'=IF(A{r}="","",G{r}+H{r})')
        assumptions.cell(r, 10, f'=IF(OR(A{r}="",I{r}=0),"",F{r}/I{r})')
        assumptions.cell(r, 11, f'=IF(A{r}="","",F{r}-I{r})')
        assumptions.cell(r, 12, f'=IF(A{r}="","",MAX(0,B{r}-H{r}))')
        assumptions.cell(r, 13, f'=IF(A{r}="","",F{r}-G{r})')
        assumptions.cell(r, 14, f'=IF(OR(A{r}="",\'Model Inputs\'!$B$9=100),"",E{r}/(1-\'Model Inputs\'!$B$9/100))')

    summary.append(["KAAMKITPRO FINANCE | PROJECT SUMMARY"])
    summary.append(["Project", "='Model Inputs'!B2", "Project Type", "='Model Inputs'!B3"])
    summary.append(["Currency", "='Model Inputs'!B4", "Projection Period (Years)", "='Model Inputs'!B13"])
    summary.append(["Project Cost", "='Model Inputs'!B5", "Debt Funding", "='Model Inputs'!B5*'Model Inputs'!B6/100"])
    summary.append(["Equity Funding", "='Model Inputs'!B5*(1-'Model Inputs'!B6/100)", "Debt / Equity", '=IF(B5=0,"",D4/B5)'])
    summary.append([])
    summary.append(["YEAR", "REVENUE", "EBITDA", "DEBT SERVICE", "DSCR", "CASH AFTER DEBT (PRE-TAX)"])
    for year in range(1, 101):
        r = 7 + year
        source = 9 + year
        summary.cell(r, 1, f'=IF(\'Assumptions\'!A{source}="","",\'Assumptions\'!A{source})')
        for col, source_col in ((2,"C"),(3,"F"),(4,"I"),(5,"J"),(6,"K")):
            summary.cell(r, col, f'=IF($A{r}="","",\'Assumptions\'!{source_col}{source})')
    summary["A110"] = "Planning note"
    summary["B110"] = "Illustrative estimate only; validate assumptions before lending or investment decisions."

    for ws in (summary, assumptions, inputs, costs, notes):
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "landscape"
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2)
        ws.oddFooter.center.text = "KaamKitPro Finance | Illustrative planning estimate"
        ws.oddFooter.right.text = "Page &P of &N"

    for ws, end_col in ((summary, 6), (assumptions, 14), (inputs, 2), (costs, 3), (notes, 1)):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
        ws["A1"].font = Font(bold=True, color=WHITE, size=16)
        ws["A1"].alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 30

    for ws, header_row in ((summary,7),(assumptions,9),(costs,2)):
        for cell in ws[header_row]:
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(bold=True, color=WHITE)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[header_row].height = 34

    for col, width in enumerate([10,18,18,18,18,18,16,16,18,12,24,18,32,22],1):
        assumptions.column_dimensions[get_column_letter(col)].width = width
    for col in "ABCDEF":
        summary.column_dimensions[col].width = 22
    summary.column_dimensions["A"].width = 25
    summary.freeze_panes = "A8"
    assumptions.freeze_panes = "A10"
    summary.auto_filter.ref = "A7:F107"
    assumptions.auto_filter.ref = "A9:N109"

    money_format = '"INR " #,##0.00;[Red]("INR " #,##0.00);-'
    for r in range(10,110):
        for col in (2,3,4,5,6,7,8,9,11,12,13,14):
            assumptions.cell(r,col).number_format = money_format
        assumptions.cell(r,10).number_format = '0.00"x"'
        for col in range(1,15):
            assumptions.cell(r,col).border = Border(bottom=THIN_GREY)
    for r in range(8,108):
        for col in (2,3,4,6):
            summary.cell(r,col).number_format = money_format
        summary.cell(r,5).number_format = '0.00"x"'
        for col in range(1,7):
            summary.cell(r,col).border = Border(bottom=THIN_GREY)
    for cell in ("B4","D4","B5"):
        summary[cell].number_format = money_format
    summary["D5"].number_format = "0.00x"

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Revenue vs EBITDA"
    chart.y_axis.title = currency
    chart.x_axis.title = "Year"
    chart.height = 7
    chart.width = 14
    chart.add_data(Reference(summary, min_col=2, max_col=3, min_row=7, max_row=107), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=8, max_row=107))
    summary.add_chart(chart, "A16")

    notes["A1"] = "Important limitations"
    notes.append(["This workbook is an illustrative planning estimate, not an audited financial statement or bank-approved DPR."])
    notes.append(["The debt schedule uses equal annual principal repayments with interest calculated on opening balance; it is not an EMI schedule."])
    notes.append(["Profit before tax shown is before depreciation and tax; it is not net profit."])
    notes.append(["Edit yellow cells on Model Inputs; the annual forecast, summary and chart recalculate in Excel."])
    notes.append(["Cash after debt service is EBITDA less interest and principal; it is pre-tax and pre-capex, not free cash flow."])
    notes.append(["Review assumptions and formulas with a qualified finance professional before submitting to a lender or investor."])
    notes.append(["PDF exports use the assumptions submitted to the API; workbook-only edits are not reflected in a PDF unless those updated assumptions are submitted again."])
    notes.column_dimensions["A"].width = 115
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE, size=14)
    for r in range(2, notes.max_row + 1):
        notes.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[r].height = 36

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def build_pdf_report(projection: dict[str, Any]) -> bytes:
    """Return a paginated PDF summary with wrapped rows and repeatable headers."""
    from textwrap import wrap

    heading = [
        "KaamKitPro Finance - Project Projection",
        f"Project: {projection.get('project_name', 'New Project')} | Type: {projection.get('project_type', 'Other')} | Currency: INR",
        f"Project cost: INR {projection.get('project_cost', 0):,.0f} | Debt: INR {projection.get('debt', 0):,.0f} | Equity: INR {projection.get('equity', 0):,.0f}",
        "Cash after debt (pre-tax, pre-capex) = EBITDA less interest and principal; not free cash flow.",
        "",
    ]
    table_header = (
        "Year | Opening debt | Revenue | Variable cost | Fixed cost | EBITDA | "
        "Interest | Principal | Closing debt | DSCR | Cash after debt (pre-tax, pre-capex) | "
        "PBT before depreciation | Break-even revenue"
    )
    rows = []
    for row in projection.get("annual_projection", []):
        dscr = "-" if row.get("dscr") is None else f"{row['dscr']:.2f}x"
        break_even = (
            f"{row['break_even_revenue']:,.0f}"
            if row.get("break_even_revenue") is not None
            else "-"
        )
        rows.append(
            f"{row.get('year', '')} | {row.get('opening_debt', 0):,.0f} | "
            f"{row.get('revenue', 0):,.0f} | {row.get('variable_cost', 0):,.0f} | "
            f"{row.get('fixed_cost', 0):,.0f} | {row.get('ebitda', 0):,.0f} | "
            f"{row.get('interest', 0):,.0f} | {row.get('principal', 0):,.0f} | "
            f"{row.get('closing_debt', 0):,.0f} | {dscr} | "
            f"{row.get('cash_after_debt_service', 0):,.0f} | "
            f"{row.get('profit_before_tax_before_depreciation', row.get('ebitda', 0) - row.get('interest', 0)):,.0f} | "
            f"{break_even}"
        )
    disclaimer = [
        "",
        "Disclaimer: Illustrative planning estimate only. Equal annual principal repayment is used, not EMI.",
        "Interest is calculated on opening balance. Profit before tax is before depreciation and tax.",
        "Cash after debt (pre-tax, pre-capex) is EBITDA less interest and principal; it is not free cash flow.",
        "This is not an audited statement, loan approval, or bank-approved DPR. Validate assumptions before use.",
    ]

    def pdf_escape(value: str) -> str:
        value = value.encode("ascii", "replace").decode("ascii")
        return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    # Wrap wide text so project details and disclaimers remain inside page margins.
    wrapped_heading = []
    for line in heading:
        wrapped_heading.extend(wrap(line, width=112) or [""])
    wrapped_rows = []
    for row in rows:
        wrapped_rows.extend(wrap(row, width=112) or [""])
    wrapped_disclaimer = []
    for line in disclaimer:
        wrapped_disclaimer.extend(wrap(line, width=112) or [""])

    # Repeat the report heading and column labels on each page.
    body_lines = wrapped_rows + wrapped_disclaimer
    first_page_capacity = 44 - len(wrapped_heading) - 2
    other_page_capacity = 44 - len(wrapped_heading) - 2
    page_chunks = []
    if not body_lines:
        page_chunks.append([])
    else:
        capacity = max(1, first_page_capacity)
        page_chunks.append(body_lines[:capacity])
        remaining = body_lines[capacity:]
        while remaining:
            page_chunks.append(remaining[:other_page_capacity])
            remaining = remaining[other_page_capacity:]

    objects: list[bytes] = []
    page_ids = []
    content_ids = []
    # Object IDs: catalog=1, pages=2, font=3, then page/content pairs.
    objects.extend([
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # Filled with the final page references below.
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ])
    for page_index, chunk in enumerate(page_chunks):
        page_id = 4 + page_index * 2
        content_id = page_id + 1
        page_ids.append(page_id)
        content_ids.append(content_id)
        lines = wrapped_heading + [table_header] + chunk
        commands = ["BT", "/F1 7 Tf", "32 560 Td", "10 TL"]
        for line_index, line in enumerate(lines):
            if line_index:
                commands.append("T*")
            commands.append(f"({pdf_escape(line)}) Tj")
        commands.append("ET")
        stream = "\n".join(commands).encode("ascii")
        page_object = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode("ascii")
        content_object = (
            b"<< /Length " + str(len(stream)).encode("ascii") +
            b" >>\nstream\n" + stream + b"\nendstream"
        )
        objects.extend([page_object, content_object])

    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("ascii")

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n% Cash after debt (pre-tax, pre-capex) disclosure\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(output)
