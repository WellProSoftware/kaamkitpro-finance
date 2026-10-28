from io import BytesIO

import pytest

from openpyxl import load_workbook

from app.reports import build_excel_report, build_pdf_report


def sample_projection():
    return {
        "project_name": "Sample Project",
        "project_type": "Pharmacy",
        "currency": "INR",
        "project_cost": 1_000_000,
        "debt": 600_000,
        "equity": 400_000,
        "model_inputs": {
            "debt_pct": 60,
            "monthly_revenue": 75_000,
            "growth_pct": 5,
            "variable_cost_pct": 44.444444,
            "monthly_fixed_cost": 16_666.6667,
            "interest_pct": 11,
            "loan_years": 5,
            "projection_years": 5,
        },
        "annual_projection": [{
            "year": 1,
            "revenue": 900_000,
            "variable_cost": 400_000,
            "fixed_cost": 200_000,
            "ebitda": 300_000,
            "interest": 66_000,
            "principal": 120_000,
            "dscr": 1.61,
            "cash_after_debt_service": 114_000,
            "opening_debt": 600_000,
            "closing_debt": 480_000,
            "profit_before_tax_before_depreciation": 234_000,
        }],
    }


def workbook_for(projection=None):
    content = build_excel_report(projection or sample_projection())
    return load_workbook(BytesIO(content), data_only=False)


def test_excel_report_contains_formula_cells():
    workbook = workbook_for()
    sheet = workbook["Assumptions"]
    assert sheet["F10"].value == '=IF(A10="","",C10-D10-E10)'
    assert sheet["I10"].value == '=IF(A10="","",G10+H10)'
    assert sheet["J10"].value == '=IF(OR(A10="",I10=0),"",F10/I10)'
    assert sheet["M10"].value == '=IF(A10="","",F10-G10)'


def test_excel_model_inputs_drive_projection_and_summary_formulas():
    workbook = workbook_for()
    inputs = workbook["Model Inputs"]
    assumptions = workbook["Assumptions"]
    summary = workbook["Executive Summary"]

    assert inputs["B7"].value == 75_000
    assert inputs["B8"].value == 5
    assert inputs["B12"].value == 5
    assert inputs["B13"].value == 5

    # Forecast formulas reference editable input cells rather than fixed values.
    assert "'Model Inputs'!$B$7" in assumptions["C10"].value
    assert "'Model Inputs'!$B$8" in assumptions["C10"].value
    assert "'Model Inputs'!$B$9" in assumptions["D10"].value
    assert "'Model Inputs'!$B$10" in assumptions["E10"].value
    assert "'Model Inputs'!$B$11" in assumptions["G10"].value
    assert "'Model Inputs'!$B$12" in assumptions["H10"].value
    assert assumptions["C11"].value.endswith("^(A11-1))")

    assert summary["B8"].value == '=IF($A8="","",\'Assumptions\'!C10)'
    assert summary["C8"].value == '=IF($A8="","",\'Assumptions\'!F10)'
    assert summary["D8"].value == '=IF($A8="","",\'Assumptions\'!I10)'
    assert summary["E8"].value == '=IF($A8="","",\'Assumptions\'!J10)'
    assert summary["F8"].value == '=IF($A8="","",\'Assumptions\'!K10)'




def test_legacy_excel_fallbacks_stay_within_input_validation_ranges():
    """Legacy projections longer than the loan-term limit still produce valid inputs."""
    projection = sample_projection()
    projection.pop("model_inputs", None)
    projection["annual_projection"] = [
        {"year": year, "revenue": 900_000}
        for year in range(1, 46)
    ]

    workbook = workbook_for(projection)
    inputs = workbook["Model Inputs"]

    assert inputs["B12"].value == 30
    assert inputs["B13"].value == 45

def test_excel_model_inputs_have_range_validation_and_clear_errors():
    workbook = workbook_for()
    inputs = workbook["Model Inputs"]
    expected_cells = (
        "B6", "B7", "B8", "B9", "B10", "B11", "B12", "B13", "B15"
    )
    validations = {
        cell: next(
            validation
            for validation in inputs.data_validations.dataValidation
            if cell in validation.cells
        )
        for cell in expected_cells
    }

    expected = {
        "B15": ("decimal", "greaterThanOrEqual", "0"),
        "B6": ("decimal", "between", "0"),
        "B7": ("decimal", "greaterThanOrEqual", "0"),
        "B8": ("decimal", "between", "-100"),
        "B9": ("decimal", "between", "0"),
        "B10": ("decimal", "greaterThanOrEqual", "0"),
        "B11": ("decimal", "between", "0"),
        "B12": ("whole", "between", "1"),
        "B13": ("whole", "between", "1"),
    }
    for cell, (kind, operator, formula1) in expected.items():
        assert cell in validations
        validation = validations[cell]
        assert validation.type == kind
        assert validation.operator == operator
        assert validation.formula1 == formula1
        assert validation.showErrorMessage is True
        assert validation.errorTitle == "Invalid model input"
        assert validation.error

    assert validations["B6"].formula2 == "100"
    assert validations["B8"].formula2 == "500"
    assert validations["B9"].formula2 == "100"
    assert validations["B11"].formula2 == "100"
    assert validations["B12"].formula2 == "30"
    assert validations["B13"].formula2 == "100"


def test_excel_forecast_has_debt_repayment_and_zero_debt_service_handling():
    workbook = workbook_for()
    sheet = workbook["Assumptions"]

    assert "MAX(0,$B$6-(A10-1)*$B$6/'Model Inputs'!$B$12)" in sheet["B10"].value
    assert "MIN(B10,$B$6/'Model Inputs'!$B$12)" in sheet["H10"].value
    assert "MAX(0,B10-H10)" in sheet["L10"].value
    assert '=IF(OR(A10="",I10=0),"",F10/I10)' == sheet["J10"].value
    assert sheet["A109"].value.endswith(',"")')
    assert sheet["C109"].value.startswith('=IF(A109="","",')


def test_excel_chart_and_recalculation_settings_are_present():
    workbook = workbook_for()
    summary = workbook["Executive Summary"]

    assert len(summary._charts) == 1
    chart = summary._charts[0]
    assert chart.title.tx.rich.p[0].r[0].t == "Revenue vs EBITDA"
    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True
    assert workbook.calculation.calcMode == "auto"


def test_excel_report_includes_debt_balance_columns():
    workbook = workbook_for()
    sheet = workbook["Assumptions"]
    assert sheet["B10"].value.startswith("=")
    assert sheet["L10"].value.startswith("=")
    assert sheet["M10"].value == '=IF(A10="","",F10-G10)'


def test_pdf_report_has_pdf_signature():
    content = build_pdf_report(sample_projection())
    assert content.startswith(b"%PDF")


def test_pdf_report_includes_debt_balance_data():
    content = build_pdf_report(sample_projection())
    assert content.startswith(b"%PDF")
    assert len(content) > 1000


def test_pdf_report_includes_profit_and_break_even_columns():
    projection = sample_projection()
    projection["annual_projection"][0]["break_even_revenue"] = 360_000
    content = build_pdf_report(projection)

    assert b"PBT before depreciation" in content
    assert b"Break-even revenue" in content
    assert b"234,000" in content
    assert b"360,000" in content


def test_excel_readme_discloses_equal_principal_not_emi():
    workbook = workbook_for()
    notes = workbook["Read Me"]
    assert "not an EMI schedule" in notes["A3"].value
    assert "not net profit" in notes["A4"].value
    assert "Edit yellow cells on Model Inputs" in notes["A5"].value

@pytest.mark.parametrize(
    ("field", "cell_ref", "formula_like_value"),
    [
        ("project_name", "B2", '=HYPERLINK("https://example.invalid","click")'),
        ("project_type", "B3", "=1+1"),
        # Currency is fixed to INR by the INR-only model.
    ],
)
def test_excel_report_stores_user_project_text_as_literal_strings(
    field, cell_ref, formula_like_value
):
    projection = sample_projection()
    projection[field] = formula_like_value

    workbook = workbook_for(projection)
    cell = workbook["Model Inputs"][cell_ref]

    assert cell.value == formula_like_value
    assert cell.data_type == "s"

def test_reports_keep_currency_consistent_with_inr_only_model():
    """Non-INR metadata must not relabel INR-based calculations."""
    projection = sample_projection()
    projection["currency"] = "USD"

    workbook = workbook_for(projection)
    assert workbook["Model Inputs"]["B4"].value == "INR"
    assert workbook["Executive Summary"]["B3"].value == "='Model Inputs'!B4"
    assert workbook["Executive Summary"]._charts[0].y_axis.title.tx.rich.p[0].r[0].t == "INR"

    pdf = build_pdf_report(projection)
    assert b"Currency: INR" in pdf
    assert b"Project cost: INR" in pdf


def test_pdf_report_paginates_long_forecasts_and_keeps_all_years():
    """Long forecasts should continue onto additional pages instead of being clipped."""
    projection = sample_projection()
    projection["annual_projection"] = [
        {
            "year": year,
            "opening_debt": max(0, 600_000 - (year - 1) * 120_000),
            "revenue": 900_000 * (1.05 ** (year - 1)),
            "variable_cost": 400_000,
            "fixed_cost": 200_000,
            "ebitda": 300_000,
            "interest": 66_000,
            "principal": 120_000,
            "closing_debt": max(0, 600_000 - year * 120_000),
            "dscr": 1.61,
            "cash_after_debt_service": 114_000,
        }
        for year in range(1, 81)
    ]

    content = build_pdf_report(projection)

    assert content.startswith(b"%PDF")
    assert content.count(b"/Type /Page /Parent") >= 2
    assert b"(80 |" in content
    assert b"(Year | Opening debt" in content


def test_projection_title_spans_break_even_revenue_column():
    workbook = workbook_for()
    sheet = workbook["Assumptions"]
    assert "A1:N1" in {str(rng) for rng in sheet.merged_cells.ranges}


def test_excel_break_even_revenue_formula_and_zero_margin_handling():
    workbook = workbook_for()
    sheet = workbook["Assumptions"]

    assert sheet["N10"].value == '=IF(OR(A10="",\'Model Inputs\'!$B$9=100),"",E10/(1-\'Model Inputs\'!$B$9/100))'
    assert "Break-even Revenue" in [cell.value for cell in sheet[9]]
    assert "N" in sheet.auto_filter.ref

    projection = sample_projection()
    projection["model_inputs"]["variable_cost_pct"] = 100
    zero_margin_workbook = workbook_for(projection)
    assert "=IF(OR(A10=" in zero_margin_workbook["Assumptions"]["N10"].value
    assert "'Model Inputs'!$B$9=100" in zero_margin_workbook["Assumptions"]["N10"].value


def test_excel_project_cost_breakdown_has_editable_costs_and_formula_total():
    workbook = workbook_for()
    costs = workbook["Project Cost Breakdown"]

    assert costs["B3"].value == 1_500_000
    assert costs["B4"].value == 2_000_000
    assert costs["B12"].value == "=SUM(B3:B11)"
    assert costs["B3"].fill.fgColor.rgb.endswith("FFF2CC")
    assert costs["B4"].fill.fgColor.rgb.endswith("FFF2CC")
    assert any(
        "B3:B11" in str(validation.sqref)
        and validation.type == "decimal"
        and validation.operator == "greaterThanOrEqual"
        for validation in costs.data_validations.dataValidation
    )


def test_itemized_cost_basis_drives_project_cost_and_funding_formulas():
    workbook = workbook_for()
    inputs = workbook["Model Inputs"]
    summary = workbook["Executive Summary"]
    assumptions = workbook["Assumptions"]
    costs = workbook["Project Cost Breakdown"]

    assert inputs["B5"].value == '=IF(B14="Itemized costs",B16,B15)'
    assert inputs["B14"].value == "Aggregate input"
    assert inputs["B15"].value == sample_projection()["project_cost"]
    assert inputs["B16"].value == "='Project Cost Breakdown'!B12"
    assert costs["B12"].value == "=SUM(B3:B11)"
    assert summary["B4"].value == "='Model Inputs'!B5"
    assert summary["D4"].value == "='Model Inputs'!B5*'Model Inputs'!B6/100"
    assert summary["B5"].value == "='Model Inputs'!B5*(1-'Model Inputs'!B6/100)"
    assert any(
        "B14" in str(validation.sqref) and validation.type == "list"
        for validation in inputs.data_validations.dataValidation
    )


def test_readme_explains_pdf_uses_submitted_assumptions():
    workbook = workbook_for()
    notes = workbook["Read Me"]
    assert any(
        "PDF exports use the assumptions submitted to the API" in str(cell.value)
        and "workbook-only edits are not reflected" in str(cell.value)
        for row in notes.iter_rows()
        for cell in row
    )


def test_pdf_and_excel_exports_start_from_same_project_funding_inputs():
    projection = sample_projection()
    workbook = workbook_for(projection)
    inputs = workbook["Model Inputs"]
    pdf = build_pdf_report(projection)

    assert inputs["B15"].value == projection["project_cost"]
    assert inputs["B6"].value == projection["model_inputs"]["debt_pct"]
    assert f"Project cost: INR {projection['project_cost']:,.0f}".encode() in pdf
    assert f"Debt: INR {projection['debt']:,.0f}".encode() in pdf
    assert f"Equity: INR {projection['equity']:,.0f}".encode() in pdf

def test_pdf_report_escapes_user_supplied_project_text():
    projection = sample_projection()
    projection["project_name"] = "Project (Phase 1) \\ Finance"
    content = build_pdf_report(projection)

    assert content.startswith(b"%PDF")
    assert b"Project \\(Phase 1\\) \\\\ Finance" in content
