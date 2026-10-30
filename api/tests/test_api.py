from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def payload():
    return {
        "project_name": "Test Pharmacy",
        "project_type": "Pharmacy",
        "project_cost": 1_000_000,
        "debt_pct": 60,
        "monthly_revenue": 150_000,
        "growth_pct": 5,
        "variable_cost_pct": 55,
        "monthly_fixed_cost": 25_000,
        "interest_pct": 11,
        "loan_years": 5,
        "projection_years": 3,
    }


def test_health_endpoint():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_root_lists_docs_and_endpoints():
    response = client.get("/")
    assert response.status_code == 200
    assert response.json()["docs"] == "/docs"
    assert "/api/reports/excel" in response.json()["endpoints"]


def test_projection_endpoint_returns_requested_years():
    response = client.post("/api/projection", json=payload())
    assert response.status_code == 200
    body = response.json()
    assert body["project_name"] == "Test Pharmacy"
    assert len(body["annual_projection"]) == 3
    assert body["debt"] + body["equity"] == body["project_cost"]


def test_excel_endpoint_returns_xlsx():
    response = client.post("/api/reports/excel", json=payload())
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert response.content.startswith(b"PK")


def test_pdf_endpoint_returns_pdf():
    response = client.post("/api/reports/pdf", json=payload())
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_rejects_invalid_debt_percentage():
    response = client.post("/api/projection", json={**payload(), "debt_pct": 150})
    assert response.status_code == 422


def test_rejects_negative_project_cost():
    response = client.post("/api/projection", json={**payload(), "project_cost": -1})
    assert response.status_code == 422


def test_rejects_unreasonably_large_project_cost():
    response = client.post(
        "/api/projection",
        json={**payload(), "project_cost": 1e16},
    )
    assert response.status_code == 422


def test_excel_endpoint_rejects_invalid_assumptions():
    response = client.post(
        "/api/reports/excel",
        json={**payload(), "variable_cost_pct": 101},
    )
    assert response.status_code == 422


def test_pdf_endpoint_rejects_invalid_assumptions():
    response = client.post(
        "/api/reports/pdf",
        json={**payload(), "loan_years": 0},
    )
    assert response.status_code == 422


def test_projection_accepts_more_than_ten_years():
    response = client.post(
        "/api/projection",
        json={**payload(), "projection_years": 25},
    )
    assert response.status_code == 200
    assert len(response.json()["annual_projection"]) == 25


def test_projection_rejects_more_than_one_hundred_years():
    response = client.post(
        "/api/projection",
        json={**payload(), "projection_years": 101},
    )
    assert response.status_code == 422


def test_cors_allows_local_frontend_origin():
    response = client.options(
        "/api/projection",
        headers={
            "Origin": "http://localhost:3000",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://localhost:3000"


def test_cors_does_not_allow_unconfigured_origin():
    response = client.options(
        "/api/projection",
        headers={
            "Origin": "https://untrusted.example",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    assert "access-control-allow-origin" not in response.headers


def test_projection_includes_reconciled_loan_schedule():
    response = client.post("/api/projection", json=payload())
    assert response.status_code == 200
    body = response.json()
    assert len(body["debt_schedule"]) == 3
    for row in body["debt_schedule"]:
        assert round(row["opening_balance"] - row["principal_repayment"], 2) == row["closing_balance"]
        assert round(row["interest_payment"] + row["principal_repayment"], 2) == row["total_debt_service"]


def test_report_endpoints_include_download_filenames():
    excel = client.post("/api/reports/excel", json=payload())
    pdf = client.post("/api/reports/pdf", json=payload())
    assert "filename=\"kaamkitpro-finance-projection.xlsx\"" in excel.headers["content-disposition"]
    assert "filename=\"kaamkitpro-finance-projection.pdf\"" in pdf.headers["content-disposition"]


def test_pdf_and_excel_endpoints_use_the_same_submitted_funding_assumptions():
    from io import BytesIO

    from openpyxl import load_workbook

    request = {**payload(), "project_cost": 3_500_000, "debt_pct": 60}
    excel_response = client.post("/api/reports/excel", json=request)
    pdf_response = client.post("/api/reports/pdf", json=request)

    assert excel_response.status_code == 200
    assert pdf_response.status_code == 200
    workbook = load_workbook(BytesIO(excel_response.content), data_only=False)
    inputs = workbook["Model Inputs"]
    pdf = pdf_response.content

    assert inputs["B15"].value == request["project_cost"]
    assert inputs["B6"].value == request["debt_pct"]
    assert f"Project cost: INR {request['project_cost']:,.0f}".encode() in pdf
    assert f"Debt: INR {request['project_cost'] * request['debt_pct'] / 100:,.0f}".encode() in pdf
    assert f"Equity: INR {request['project_cost'] * (1 - request['debt_pct'] / 100):,.0f}".encode() in pdf


def test_itemized_costs_drive_projection_and_both_reports():
    from io import BytesIO

    from openpyxl import load_workbook

    itemized = {
        "Outsourcing": 1_500_000,
        "Warehouse setup": 2_000_000,
        "Hospital setup / expansion": 500_000,
        "Retail / pharmacy setup or expansion": 0,
        "Optical setup or expansion": 0,
        "Equipment and technology": 0,
        "Licences, deposits and professional fees": 0,
        "Initial inventory and working capital": 0,
        "Other project costs": 0,
    }
    request = {**payload(), "project_cost": 1_000_000, "debt_pct": 60, "itemized_costs": itemized}
    projection = client.post("/api/projection", json=request)
    excel = client.post("/api/reports/excel", json=request)
    pdf = client.post("/api/reports/pdf", json=request)

    assert projection.status_code == 200
    assert projection.json()["project_cost"] == 4_000_000
    assert projection.json()["debt"] == 2_400_000
    assert projection.json()["equity"] == 1_600_000
    assert excel.status_code == 200 and pdf.status_code == 200

    workbook = load_workbook(BytesIO(excel.content), data_only=False)
    inputs = workbook["Model Inputs"]
    costs = workbook["Project Cost Breakdown"]
    assert inputs["B14"].value == "Itemized costs"
    assert inputs["B5"].value == '=IF(B14="Itemized costs",B16,B15)'
    assert costs["B3"].value == 1_500_000
    assert costs["B4"].value == 2_000_000
    assert costs["B5"].value == 500_000
    assert costs["B12"].value == "=SUM(B3:B11)"
    assert b"Project cost: INR 4,000,000" in pdf.content
    assert b"Debt: INR 2,400,000" in pdf.content
    assert b"Equity: INR 1,600,000" in pdf.content


def test_itemized_costs_reject_unknown_categories_and_negative_amounts():
    unknown = client.post(
        "/api/projection",
        json={**payload(), "itemized_costs": {"Mystery cost": 100}},
    )
    negative = client.post(
        "/api/projection",
        json={**payload(), "itemized_costs": {"Outsourcing": -1}},
    )
    oversized_total = client.post(
        "/api/projection",
        json={**payload(), "itemized_costs": {"Outsourcing": 1_000_000_000_000_001}},
    )
    assert unknown.status_code == 422
    assert negative.status_code == 422
    assert oversized_total.status_code == 422

def test_projection_handles_zero_debt_without_division_errors():
    response = client.post(
        "/api/projection",
        json={**payload(), "project_cost": 0, "debt_pct": 0},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["debt"] == 0
    assert body["equity"] == 0
    assert all(row["debt_service"] == 0 for row in body["annual_projection"])
    assert all(row["dscr"] is None for row in body["annual_projection"])


def test_projection_handles_one_hundred_percent_variable_cost():
    response = client.post(
        "/api/projection",
        json={**payload(), "variable_cost_pct": 100},
    )
    assert response.status_code == 200
    body = response.json()
    assert all(row["break_even_revenue"] is None for row in body["annual_projection"])
    assert all(row["variable_cost"] == row["revenue"] for row in body["annual_projection"])

def test_loan_is_fully_repaid_after_term_and_interest_stops():
    response = client.post(
        "/api/projection",
        json={**payload(), "project_cost": 1_000_000, "debt_pct": 60, "loan_years": 5, "projection_years": 7},
    )
    assert response.status_code == 200
    body = response.json()
    schedule = body["debt_schedule"]

    assert len(schedule) == 7
    assert schedule[4]["closing_balance"] == 0
    assert schedule[5]["opening_balance"] == 0
    assert schedule[5]["principal_repayment"] == 0
    assert schedule[5]["interest_payment"] == 0
    assert schedule[5]["total_debt_service"] == 0
    assert schedule[6]["closing_balance"] == 0

def test_projection_supports_zero_and_full_debt_funding_boundaries():
    no_debt = client.post(
        "/api/projection",
        json={**payload(), "project_cost": 1_000_000, "debt_pct": 0},
    )
    full_debt = client.post(
        "/api/projection",
        json={**payload(), "project_cost": 1_000_000, "debt_pct": 100},
    )

    assert no_debt.status_code == 200
    assert no_debt.json()["debt"] == 0
    assert no_debt.json()["equity"] == 1_000_000
    assert full_debt.status_code == 200
    assert full_debt.json()["debt"] == 1_000_000
    assert full_debt.json()["equity"] == 0


def test_projection_supports_zero_interest_rate():
    response = client.post(
        "/api/projection",
        json={**payload(), "interest_pct": 0},
    )

    assert response.status_code == 200
    body = response.json()
    assert all(row["interest"] == 0 for row in body["annual_projection"])
    assert all(
        row["debt_service"] == row["principal"]
        for row in body["annual_projection"]
    )

def test_projection_handles_total_revenue_decline_after_year_one():
    response = client.post(
        "/api/projection",
        json={**payload(), "growth_pct": -100, "projection_years": 3},
    )

    assert response.status_code == 200
    rows = response.json()["annual_projection"]
    assert rows[0]["revenue"] == payload()["monthly_revenue"] * 12
    assert rows[1]["revenue"] == 0
    assert rows[2]["revenue"] == 0
    assert all(row["variable_cost"] == 0 for row in rows[1:])


def test_projection_applies_maximum_supported_growth_rate_consistently():
    response = client.post(
        "/api/projection",
        json={**payload(), "growth_pct": 500, "projection_years": 3},
    )

    assert response.status_code == 200
    rows = response.json()["annual_projection"]
    first_year_revenue = payload()["monthly_revenue"] * 12
    assert rows[0]["revenue"] == first_year_revenue
    assert rows[1]["revenue"] == first_year_revenue * 6
    assert rows[2]["revenue"] == first_year_revenue * 36

def test_projection_keeps_fixed_costs_when_revenue_falls_to_zero():
    response = client.post(
        "/api/projection",
        json={
            **payload(),
            "growth_pct": -100,
            "monthly_fixed_cost": 25_000,
            "projection_years": 2,
        },
    )

    assert response.status_code == 200
    year_two = response.json()["annual_projection"][1]
    assert year_two["revenue"] == 0
    assert year_two["variable_cost"] == 0
    assert year_two["fixed_cost"] == 300_000
    assert year_two["ebitda"] == -300_000
    assert year_two["cash_after_debt_service"] < year_two["ebitda"]

def test_projection_keeps_revenue_constant_with_zero_growth():
    response = client.post(
        "/api/projection",
        json={**payload(), "growth_pct": 0, "projection_years": 5},
    )

    assert response.status_code == 200
    rows = response.json()["annual_projection"]
    expected_annual_revenue = payload()["monthly_revenue"] * 12
    assert len(rows) == 5
    assert all(row["revenue"] == expected_annual_revenue for row in rows)
    assert all(
        row["variable_cost"] == expected_annual_revenue * payload()["variable_cost_pct"] / 100
        for row in rows
    )

def test_projection_applies_negative_growth_as_compound_decline():
    response = client.post(
        "/api/projection",
        json={**payload(), "growth_pct": -50, "projection_years": 3},
    )

    assert response.status_code == 200
    rows = response.json()["annual_projection"]
    first_year_revenue = payload()["monthly_revenue"] * 12
    assert [row["revenue"] for row in rows] == [
        first_year_revenue,
        first_year_revenue * 0.5,
        first_year_revenue * 0.25,
    ]

def test_projection_break_even_equals_fixed_cost_when_variable_cost_is_zero():
    response = client.post(
        "/api/projection",
        json={
            **payload(),
            "variable_cost_pct": 0,
            "monthly_fixed_cost": 40_000,
            "projection_years": 2,
        },
    )

    assert response.status_code == 200
    rows = response.json()["annual_projection"]
    assert all(row["variable_cost"] == 0 for row in rows)
    assert all(row["break_even_revenue"] == 480_000 for row in rows)

def test_projection_rejects_fractional_loan_and_projection_years():
    fractional_loan = client.post(
        "/api/projection",
        json={**payload(), "loan_years": 2.5},
    )
    fractional_projection = client.post(
        "/api/projection",
        json={**payload(), "projection_years": 3.5},
    )

    assert fractional_loan.status_code == 422
    assert fractional_projection.status_code == 422

def test_projection_rejects_project_labels_over_field_limits():
    long_name = client.post(
        "/api/projection",
        json={**payload(), "project_name": "P" * 121},
    )
    long_type = client.post(
        "/api/projection",
        json={**payload(), "project_type": "T" * 81},
    )

    assert long_name.status_code == 422
    assert long_type.status_code == 422


def test_projection_accepts_project_labels_at_field_limits():
    response = client.post(
        "/api/projection",
        json={**payload(), "project_name": "P" * 120, "project_type": "T" * 80},
    )

    assert response.status_code == 200
    assert len(response.json()["project_name"]) == 120
    assert len(response.json()["project_type"]) == 80

def test_excel_report_has_formula_recalculation_and_input_validations():
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.post("/api/reports/excel", json=payload())
    assert response.status_code == 200

    workbook = load_workbook(BytesIO(response.content), data_only=False)
    inputs = workbook["Model Inputs"]
    assumptions = workbook["Assumptions"]
    validations = {
        validation.sqref.__str__(): validation.type
        for validation in inputs.data_validations.dataValidation
    }

    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True
    assert workbook.calculation.calcMode == "auto"
    assert "B12" in validations and validations["B12"] == "whole"
    assert "B13" in validations and validations["B13"] == "whole"
    assert assumptions["C10"].value.startswith("=IF(")
    assert assumptions["F10"].value.startswith("=IF(")
    assert assumptions["N10"].value.startswith("=IF(")

def test_long_projection_pdf_is_paginated_and_repeats_report_headers():
    response = client.post(
        "/api/reports/pdf",
        json={**payload(), "projection_years": 100},
    )

    assert response.status_code == 200
    pdf = response.content
    assert b"/Count " in pdf
    page_count = int(pdf.split(b"/Count ", 1)[1].split(b" ", 1)[0])
    assert page_count > 1
    assert pdf.count(b"Year | Opening debt | Revenue") == page_count
    assert b"This is not an audited statement" in pdf

def test_projection_cash_after_debt_service_is_ebitda_less_interest_and_principal():
    response = client.post("/api/projection", json=payload())

    assert response.status_code == 200
    for row in response.json()["annual_projection"]:
        expected = row["ebitda"] - row["interest"] - row["principal"]
        assert row["cash_after_debt_service"] == expected
        assert row["debt_service"] == row["interest"] + row["principal"]

def test_reports_label_cash_after_debt_service_as_pre_tax_and_pre_capex():
    from io import BytesIO

    from openpyxl import load_workbook

    excel = client.post("/api/reports/excel", json=payload())
    pdf = client.post("/api/reports/pdf", json=payload())

    assert excel.status_code == 200
    assert pdf.status_code == 200
    workbook = load_workbook(BytesIO(excel.content), data_only=False)
    assumptions = workbook["Assumptions"]
    summary = workbook["Executive Summary"]
    notes = workbook["Read Me"]

    assert assumptions["K9"].value == "Cash After Debt Service (Pre-Tax, Pre-Capex)"
    assert summary["F7"].value == "CASH AFTER DEBT (PRE-TAX)"
    assert any("not free cash flow" in str(row[0].value) for row in notes.iter_rows())
    assert b"Cash after debt (pre-tax, pre-capex)" in pdf.content


def test_itemized_costs_reject_non_finite_and_over_limit_values():
    non_finite = client.post(
        "/api/projection",
        content='{"project_cost":1000000,"itemized_costs":{"Outsourcing":Infinity}}',
        headers={"Content-Type": "application/json"},
    )
    over_limit = client.post(
        "/api/projection",
        json={**payload(), "itemized_costs": {"Outsourcing": 1_000_000_000_000_001}},
    )
    assert non_finite.status_code == 422
    assert over_limit.status_code == 422


def test_itemized_costs_reject_sum_over_supported_limit():
    response = client.post(
        "/api/projection",
        json={**payload(), "itemized_costs": {
            "Outsourcing": 600_000_000_000_000,
            "Warehouse setup": 500_000_000_000_000,
        }},
    )
    assert response.status_code == 422


def test_excel_and_pdf_annual_rows_reconcile_to_projection_api():
    from io import BytesIO

    from openpyxl import load_workbook

    request = {
        **payload(),
        "project_cost": 4_250_000,
        "debt_pct": 65,
        "monthly_revenue": 275_000,
        "growth_pct": 12,
        "variable_cost_pct": 42,
        "monthly_fixed_cost": 38_000,
        "interest_pct": 10.5,
        "loan_years": 2,
        "projection_years": 3,
    }
    projection_response = client.post("/api/projection", json=request)
    excel_response = client.post("/api/reports/excel", json=request)
    pdf_response = client.post("/api/reports/pdf", json=request)

    assert projection_response.status_code == 200
    assert excel_response.status_code == 200
    assert pdf_response.status_code == 200

    projection = projection_response.json()
    workbook = load_workbook(BytesIO(excel_response.content), data_only=False)
    assumptions = workbook["Assumptions"]
    summary = workbook["Executive Summary"]
    pdf_text = pdf_response.content.decode("latin-1")

    for year, result in enumerate(projection["annual_projection"], start=1):
        row = 9 + year
        summary_row = 7 + year
        assert assumptions.cell(row, 1).value.startswith("=IF(ROW()-9<=")
        assert assumptions.cell(row, 3).value.startswith("=IF(A")
        assert assumptions.cell(row, 6).value.startswith("=IF(A")
        assert assumptions.cell(row, 9).value.startswith("=IF(A")
        assert assumptions.cell(row, 11).value.startswith("=IF(A")
        assert summary.cell(summary_row, 2).value.startswith("=IF($A")
        assert summary.cell(summary_row, 6).value.startswith("=IF($A")

        expected_pdf_values = [
            str(year),
            f'{result["opening_debt"]:,.0f}',
            f'{result["revenue"]:,.0f}',
            f'{result["variable_cost"]:,.0f}',
            f'{result["fixed_cost"]:,.0f}',
            f'{result["ebitda"]:,.0f}',
            f'{result["interest"]:,.0f}',
            f'{result["principal"]:,.0f}',
            f'{result["closing_debt"]:,.0f}',
            "-" if result["dscr"] is None else f'{result["dscr"]:.2f}x',
            f'{result["cash_after_debt_service"]:,.0f}',
            f'{result["profit_before_tax_before_depreciation"]:,.0f}',
            "-" if result["break_even_revenue"] is None else f'{result["break_even_revenue"]:,.0f}',
        ]
        # PDF wraps long table rows across text lines to stay within page margins.
        # Decode PDF text operators and compare the rendered table row as normalized text.
        import re

        pdf_literals = re.findall(r"\\((.*?)(?<!\\)\\) Tj", pdf_text, flags=re.DOTALL)
        rendered_lines = [re.sub(r"\\s+", " ", value) for value in pdf_literals]
        normalized_values = [re.sub(r"\\s+", " ", value) for value in expected_pdf_values]
        assert all(
            any(value in line for line in rendered_lines)
            for value in normalized_values
        )

