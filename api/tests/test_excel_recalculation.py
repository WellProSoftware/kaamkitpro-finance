"""Integration test for real Excel formula calculation via LibreOffice Calc."""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.reports import build_excel_report


def recalculate_with_libreoffice(source: Path, output_dir: Path, profile: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [
            "soffice",
            "--headless",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_dir),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=90,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    calculated_path = output_dir / source.name
    assert calculated_path.exists(), result.stdout + result.stderr
    return calculated_path


def assert_no_formula_errors(workbook) -> None:
    for sheet_name, cell_range in (
        # Scan the complete supported forecast area, not only the first few years.
        ("Assumptions", ("A10", "M109")),
        ("Executive Summary", ("A8", "F107")),
    ):
        sheet = workbook[sheet_name]
        start, end = cell_range
        for row in sheet[start:end]:
            for cell in row:
                assert not (
                    isinstance(cell.value, str)
                    and cell.value.startswith("#")
                ), f"Formula error at {sheet_name}!{cell.coordinate}: {cell.value}"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
def test_excel_workbook_recalculates_live_forecast_and_summary(tmp_path: Path) -> None:
    projection = {
        "project_name": "Calc integration test",
        "project_type": "Other",
        "currency": "INR",
        "project_cost": 1_000_000,
        "debt": 600_000,
        "equity": 400_000,
        "model_inputs": {
            "debt_pct": 60,
            "monthly_revenue": 100_000,
            "growth_pct": 10,
            "variable_cost_pct": 40,
            "monthly_fixed_cost": 20_000,
            "interest_pct": 10,
            "loan_years": 5,
            "projection_years": 3,
        },
        "annual_projection": [],
    }

    source = tmp_path / "finance-model.xlsx"
    source.write_bytes(build_excel_report(projection))

    # First pass: calculate the workbook with its original assumptions.
    baseline_path = recalculate_with_libreoffice(
        source, tmp_path / "baseline", tmp_path / "lo-profile-baseline"
    )
    baseline = load_workbook(baseline_path, data_only=True)
    assumptions = baseline["Assumptions"]
    summary = baseline["Executive Summary"]

    # Revenue is 12 * monthly revenue, growing by 10% per year.
    assert assumptions["C10"].value == pytest.approx(1_200_000)
    assert assumptions["C11"].value == pytest.approx(1_320_000)
    assert assumptions["C12"].value == pytest.approx(1_452_000)

    # Debt schedule, debt-service coverage and summary references recalculate.
    assert assumptions["B10"].value == pytest.approx(600_000)
    assert assumptions["H10"].value == pytest.approx(120_000)
    assert assumptions["G10"].value == pytest.approx(60_000)
    expected_dscr = (1_200_000 - 480_000 - 240_000) / (60_000 + 120_000)
    assert assumptions["J10"].value == pytest.approx(expected_dscr)
    assert assumptions["L10"].value == pytest.approx(480_000)
    assert summary["B8"].value == pytest.approx(1_200_000)
    assert summary["C8"].value == pytest.approx(480_000)
    assert_no_formula_errors(baseline)

    # Change a yellow editable input and recalculate a second copy. The
    # formula outputs must follow the input rather than retain stale values.
    edited_source = tmp_path / "finance-model-edited.xlsx"
    editable = load_workbook(source, data_only=False)
    editable["Model Inputs"]["B7"] = 120_000
    editable.save(edited_source)

    edited_path = recalculate_with_libreoffice(
        edited_source, tmp_path / "edited", tmp_path / "lo-profile-edited"
    )
    edited = load_workbook(edited_path, data_only=True)
    edited_assumptions = edited["Assumptions"]
    edited_summary = edited["Executive Summary"]

    assert edited_assumptions["C10"].value == pytest.approx(1_440_000)
    assert edited_assumptions["C11"].value == pytest.approx(1_584_000)
    assert edited_assumptions["F10"].value == pytest.approx(624_000)
    assert edited_summary["B8"].value == pytest.approx(1_440_000)
    assert edited_summary["C8"].value == pytest.approx(624_000)
    assert edited_assumptions["C10"].value != assumptions["C10"].value
    assert_no_formula_errors(edited)


def test_excel_model_inputs_have_validation_for_all_numeric_assumptions() -> None:
    """Guard against accidentally removing spreadsheet input safeguards."""
    workbook_bytes = build_excel_report(
        {
            "project_name": "Validation test",
            "project_type": "Other",
            "currency": "INR",
            "project_cost": 1_000_000,
            "debt": 600_000,
            "annual_projection": [],
            "model_inputs": {
                "debt_pct": 60,
                "monthly_revenue": 100_000,
                "growth_pct": 10,
                "variable_cost_pct": 40,
                "monthly_fixed_cost": 20_000,
                "interest_pct": 10,
                "loan_years": 5,
                "projection_years": 3,
            },
        }
    )
    import io

    workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False)
    inputs = workbook["Model Inputs"]
    validated_coordinates = {
        coordinate
        for validation in inputs.data_validations.dataValidation
        for coordinate in str(validation.sqref).split()
    }

    # Editable numeric assumptions must have explicit validation rules.
    assert {f"B{row}" for row in range(6, 14)} | {"B15"} <= validated_coordinates

    validation_by_cell = {
        coordinate: validation
        for validation in inputs.data_validations.dataValidation
        for coordinate in str(validation.sqref).split()
    }
    assert validation_by_cell["B6"].type == "decimal"
    assert validation_by_cell["B6"].operator == "between"
    assert validation_by_cell["B12"].type == "whole"
    assert validation_by_cell["B13"].type == "whole"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
@pytest.mark.parametrize(
    ("project_cost", "debt_pct"),
    [(0, 50), (1_000_000, 0)],
    ids=("zero-project-cost", "zero-debt"),
)
def test_excel_workbook_handles_zero_cost_or_zero_debt(
    tmp_path: Path, project_cost: float, debt_pct: float
) -> None:
    """Zero-cost and zero-debt scenarios should not produce formula errors."""
    projection = {
        "project_name": "Zero-value edge case",
        "project_type": "Other",
        "currency": "INR",
        "project_cost": project_cost,
        "debt": project_cost * debt_pct / 100,
        "equity": project_cost * (1 - debt_pct / 100),
        "model_inputs": {
            "debt_pct": debt_pct,
            "monthly_revenue": 0,
            "growth_pct": 0,
            "variable_cost_pct": 40,
            "monthly_fixed_cost": 0,
            "interest_pct": 10,
            "loan_years": 5,
            "projection_years": 1,
        },
        "annual_projection": [],
    }
    source = tmp_path / "zero-case.xlsx"
    source.write_bytes(build_excel_report(projection))
    calculated_path = recalculate_with_libreoffice(
        source, tmp_path / "calculated", tmp_path / "lo-profile"
    )
    workbook = load_workbook(calculated_path, data_only=True)
    assumptions = workbook["Assumptions"]
    summary = workbook["Executive Summary"]

    assert assumptions["B10"].value == pytest.approx(project_cost * debt_pct / 100)
    assert assumptions["C10"].value == pytest.approx(0)
    assert assumptions["G10"].value == pytest.approx(0)
    assert assumptions["H10"].value == pytest.approx(0)
    assert assumptions["I10"].value == pytest.approx(0)
    assert assumptions["J10"].value in (None, "")
    if project_cost == 0:
        assert summary["D5"].value in (None, "")
    else:
        assert summary["D5"].value == pytest.approx(0)
    assert_no_formula_errors(workbook)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
def test_excel_workbook_supports_full_100_year_projection(tmp_path: Path) -> None:
    """The maximum supported horizon should populate year 100 and no later rows."""
    projection = {
        "project_name": "Maximum horizon test",
        "project_type": "Other",
        "currency": "INR",
        "project_cost": 1_000_000,
        "debt": 1_000_000,
        "equity": 0,
        "model_inputs": {
            "debt_pct": 100,
            "monthly_revenue": 10_000,
            "growth_pct": 0,
            "variable_cost_pct": 50,
            "monthly_fixed_cost": 1_000,
            "interest_pct": 0,
            "loan_years": 30,
            "projection_years": 100,
        },
        "annual_projection": [],
    }
    source = tmp_path / "maximum-horizon.xlsx"
    source.write_bytes(build_excel_report(projection))
    calculated_path = recalculate_with_libreoffice(
        source, tmp_path / "calculated", tmp_path / "lo-profile"
    )
    workbook = load_workbook(calculated_path, data_only=True)
    assumptions = workbook["Assumptions"]
    summary = workbook["Executive Summary"]

    assert assumptions["A10"].value == 1
    assert assumptions["A39"].value == 30
    assert assumptions["A109"].value == 100
    assert assumptions["C109"].value == pytest.approx(120_000)
    assert summary["A107"].value == 100
    assert summary["B107"].value == pytest.approx(120_000)
    assert assumptions["A110"].value in (None, "")
    assert_no_formula_errors(workbook)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
def test_excel_loan_is_repaid_at_term_and_interest_stops(tmp_path: Path) -> None:
    """Equal-principal debt reaches zero at term end; later interest is zero."""
    projection = {
        "project_name": "Loan payoff test",
        "project_type": "Other",
        "currency": "INR",
        "project_cost": 1_000_000,
        "debt": 600_000,
        "equity": 400_000,
        "model_inputs": {
            "debt_pct": 60,
            "monthly_revenue": 100_000,
            "growth_pct": 0,
            "variable_cost_pct": 40,
            "monthly_fixed_cost": 20_000,
            "interest_pct": 10,
            "loan_years": 5,
            "projection_years": 7,
        },
        "annual_projection": [],
    }
    source = tmp_path / "loan-payoff.xlsx"
    source.write_bytes(build_excel_report(projection))
    calculated_path = recalculate_with_libreoffice(
        source, tmp_path / "calculated", tmp_path / "lo-profile"
    )
    workbook = load_workbook(calculated_path, data_only=True)
    assumptions = workbook["Assumptions"]

    # Year 5 starts with the final 120,000 principal balance.
    assert assumptions["A14"].value == 5
    assert assumptions["B14"].value == pytest.approx(120_000)
    assert assumptions["G14"].value == pytest.approx(12_000)
    assert assumptions["H14"].value == pytest.approx(120_000)
    assert assumptions["L14"].value == pytest.approx(0)

    # Debt service contains no interest or principal after payoff.
    assert assumptions["A15"].value == 6
    assert assumptions["G15"].value == pytest.approx(0)
    assert assumptions["H15"].value == pytest.approx(0)
    assert assumptions["I15"].value == pytest.approx(0)
    assert assumptions["L15"].value == pytest.approx(0)
    assert_no_formula_errors(workbook)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
def test_excel_forecast_matches_api_model_calculations(tmp_path: Path) -> None:
    """The editable workbook and API model should agree for every forecast year."""
    from app.model import build_projection

    inputs = {
        "project_name": "Cross-check",
        "project_type": "Other",
        "project_cost": 2_500_000,
        "debt_pct": 65,
        "monthly_revenue": 180_000,
        "growth_pct": 8,
        "variable_cost_pct": 42,
        "monthly_fixed_cost": 35_000,
        "interest_pct": 9.5,
        "loan_years": 7,
        "projection_years": 12,
    }
    projection = build_projection(inputs)
    source = tmp_path / "cross-check.xlsx"
    source.write_bytes(build_excel_report({
        **projection,
        "project_name": inputs["project_name"],
        "project_type": inputs["project_type"],
    }))
    calculated_path = recalculate_with_libreoffice(
        source, tmp_path / "calculated", tmp_path / "lo-profile"
    )
    workbook = load_workbook(calculated_path, data_only=True)
    assumptions = workbook["Assumptions"]

    # Workbook rows 10 onward map to API forecast rows starting at year 1.
    fields = {
        "opening_debt": "B",
        "revenue": "C",
        "variable_cost": "D",
        "fixed_cost": "E",
        "ebitda": "F",
        "interest": "G",
        "principal": "H",
        "debt_service": "I",
        "closing_debt": "L",
        "profit_before_tax_before_depreciation": "M",
        "break_even_revenue": "N",
        "cash_after_debt_service": "K",
    }
    for index, expected in enumerate(projection["annual_projection"], start=10):
        assert assumptions[f"A{index}"].value == expected["year"]
        for field, column in fields.items():
            assert assumptions[f"{column}{index}"].value == pytest.approx(
                expected[field], abs=0.02
            )
        expected_dscr = expected["dscr"]
        actual_dscr = assumptions[f"J{index}"].value
        if expected_dscr is None:
            assert actual_dscr in (None, "")
        else:
            assert actual_dscr == pytest.approx(expected_dscr, abs=0.01)

    assert_no_formula_errors(workbook)
